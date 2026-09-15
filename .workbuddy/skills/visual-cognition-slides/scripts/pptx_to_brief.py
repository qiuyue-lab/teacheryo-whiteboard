#!/usr/bin/env python3
"""
pptx_to_brief.py — 从 PPTX 提取内容简报，供 Claude + Visual Cognition Slides Skill 重新设计

用法：
  python3 pptx_to_brief.py 课件.pptx
  python3 pptx_to_brief.py 课件.pptx -o brief.md

输出：
  一份 Markdown 简报，包含每张 slide 的标题、文字、备注、图片数。
  把这份简报粘贴给 Claude，加上以下 prompt，即可触发 skill 重新设计：

  > 这是一份已有 PPT 的内容，请用 Visual Cognition Slides skill
  > 重新设计成双通道动画 HTML slides。受众是[X]，核心目标是[Y]。

依赖：
  pip install python-pptx
"""

import sys
import re
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    print("请先安装依赖：pip install python-pptx")
    sys.exit(1)


def extract_brief(pptx_path: Path) -> str:
    prs = Presentation(str(pptx_path))
    total = len(prs.slides)
    lines = []

    lines.append(f"# 课件内容简报：{pptx_path.stem}")
    lines.append(f"\n共 {total} 张 slides\n")
    lines.append("---\n")
    lines.append(
        "> 用途：把这份简报粘贴给 Claude，告知受众和教学目标，\n"
        "> Claude 会用 Visual Cognition Slides skill 重新设计成动画 HTML。\n"
    )
    lines.append("---\n")

    for i, slide in enumerate(prs.slides):
        slide_num = i + 1
        title_text = None
        body_items = []
        image_count = 0
        notes_text = None

        # 演讲者备注
        try:
            if slide.has_notes_slide:
                raw = slide.notes_slide.notes_text_frame.text.strip()
                if raw:
                    notes_text = raw
        except Exception:
            pass

        # 形状
        for shape in slide.shapes:
            # 图片计数
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image_count += 1

            if not shape.has_text_frame:
                continue

            for para in shape.text_frame.paragraphs:
                text = para.text.strip()
                if not text:
                    continue

                # 判断是否标题（利用占位符类型或字号）
                is_title = False
                try:
                    from pptx.enum.text import PP_PLACEHOLDER
                    if shape.is_placeholder and shape.placeholder_format.type in (
                        PP_PLACEHOLDER.TITLE,
                        PP_PLACEHOLDER.CENTER_TITLE,
                    ):
                        is_title = True
                except Exception:
                    pass

                if not is_title and para.runs:
                    try:
                        size = para.runs[0].font.size
                        if size and size / 12700 > 28 and not title_text:
                            is_title = True
                    except Exception:
                        pass

                if is_title and not title_text:
                    title_text = text
                else:
                    body_items.append(text)

        # 去重并过滤过短片段
        seen = set()
        body_unique = []
        for item in body_items:
            clean = re.sub(r'\s+', ' ', item)
            if clean not in seen and len(clean) > 1:
                seen.add(clean)
                body_unique.append(clean)

        # 写入简报
        heading = title_text or f"（无标题）"
        lines.append(f"## Slide {slide_num}：{heading}\n")

        if body_unique:
            lines.append("**文字内容：**")
            for item in body_unique:
                lines.append(f"- {item}")
            lines.append("")

        if image_count:
            lines.append(f"**图片**：{image_count} 张（位置与内容见原 PPT）\n")

        if notes_text:
            lines.append(f"**备注（口述稿）：**")
            # 按行缩进
            for note_line in notes_text.splitlines():
                note_line = note_line.strip()
                if note_line:
                    lines.append(f"  {note_line}")
            lines.append("")

        lines.append("---\n")

    # 结尾 prompt 建议
    lines.append("## 如何使用这份简报\n")
    lines.append(
        "把以上内容粘贴给 Claude，附加以下信息：\n\n"
        "```\n"
        "受众：[例：初中生 / 大学生 / 企业培训]\n"
        "核心目标：[学完这节课，他们能做/理解什么]\n"
        "视觉风格偏好（可选）：[手绘创意 / 学术简约 / 暖色插画 / 科研极客…]\n"
        "画布格式（可选）：[横版1920×1080 / 竖版手机 / 小红书4:5]\n"
        "```\n\n"
        "然后说：**「请用 Visual Cognition Slides skill 重新设计这份课件，"
        "加上双通道动画，生成 HTML。」**\n"
    )

    return "\n".join(lines)


def main():
    args = sys.argv[1:]
    if not args or args[0] in ("-h", "--help"):
        print(__doc__)
        sys.exit(0)

    pptx_path = Path(args[0])
    if not pptx_path.exists():
        print(f"文件不存在：{pptx_path}")
        sys.exit(1)

    output_path = None
    for i, a in enumerate(args):
        if a in ("-o", "--output") and i + 1 < len(args):
            output_path = Path(args[i + 1])

    if output_path is None:
        output_path = pptx_path.with_name(pptx_path.stem + "_brief.md")

    print(f"📂 读取：{pptx_path}")
    brief = extract_brief(pptx_path)

    output_path.write_text(brief, encoding="utf-8")
    print(f"✅ 简报已保存：{output_path}")
    print()
    print("下一步：把这份 Markdown 文件粘贴给 Claude，")
    print("告知受众和教学目标，Claude 会生成双通道动画 HTML slides。")


if __name__ == "__main__":
    main()
